"""Export tdd3.md to matching editable PowerPoint and offline HTML decks.

Install: python -m pip install -r requirements.txt
Run:     python export_tdd.py
"""

from __future__ import annotations

import html
import json
import re
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "tdd3.md"
BG, PANEL, FG, MUTED, BLUE = "0F1115", "151922", "E6E6E6", "AAAAAA", "4AA3FF"


@dataclass
class Slide:
    title: str
    section: str
    source: str
    blocks: list[str]
    diagram: bool = False


def clean(text: str) -> str:
    return re.sub(r"\*\*|~~|`", "", text)


def make_blocks(lines: list[str]) -> list[str]:
    blocks: list[str] = []
    current: list[str] = []
    fenced = False
    for line in lines:
        if line.startswith("```"):
            fenced = not fenced
        new_item = bool(re.match(r"^(?:- |\d+\. )", line))
        if not fenced and (not line.strip() or new_item) and current:
            blocks.append("\n".join(current))
            current = []
        if line.strip() or fenced:
            current.append(line)
    if current:
        blocks.append("\n".join(current))
    return blocks


def read_slides() -> tuple[str, list[Slide]]:
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    title = lines[0].removeprefix("# ")
    slides = [Slide(title, "Presentation", "\n".join(lines[:4]), [lines[2]])]
    headings: dict[int, str] = {}
    heading = ""
    section = ""
    body: list[str] = []

    def flush() -> None:
        if not any(line.strip() for line in body):
            return
        blocks = make_blocks(body)
        if any("```mermaid" in block for block in blocks):
            slides.append(Slide(heading, section, "\n".join(body), blocks, True))
            return
        chunks: list[list[str]] = [[]]
        size = 0
        for block in blocks:
            cost = len(block) + 65
            if chunks[-1] and (size + cost > 850 or len(chunks[-1]) >= 6):
                chunks.append([])
                size = 0
            chunks[-1].append(block)
            size += cost
        for n, chunk in enumerate(chunks, 1):
            suffix = f" ({n}/{len(chunks)})" if len(chunks) > 1 else ""
            slides.append(Slide(heading + suffix, section, "\n\n".join(chunk), chunk))

    for line in lines[4:]:
        match = re.match(r"^(#{2,4}) (.+)$", line)
        phase = re.match(r"^(RED|GREEN|REFACTOR) - (.+)$", line)
        purposes = line.startswith("Different phases have different purposes:")
        if match or phase or purposes:
            flush()
            body = []
            if match:
                level = len(match[1])
                heading = match[2]
                headings = {k: v for k, v in headings.items() if k < level}
                headings[level] = heading
                section = " / ".join(v for k, v in headings.items() if k < level)
                if not section:
                    section = "Appendix · Study" if heading in ("Study", "Resources", "Ideas") else "Presentation"
            elif phase:
                heading = phase[1]
                section = "4. RED–GREEN–REFACTOR"
                body.append(phase[2])
            else:
                heading = "Different phases have different purposes"
                section = "4. RED–GREEN–REFACTOR"
        else:
            body.append(line)
    flush()
    return title, slides


def link_parts(text: str) -> tuple[str, str | None]:
    match = re.search(r"https?://\S+", text)
    if not match:
        return text, None
    url = match[0]
    label = text[:match.start()].strip().removesuffix("-").strip()
    if not label:
        labels = {
            "blog.stevensanderson.com": "Steve Sanderson — Writing Great Unit Tests",
            "www.google.com": "Search: TDD, Where Did It All Go Wrong?",
            "www.reddit.com": "Discussion: Ian Cooper — TDD, Where Did It All Go Wrong?",
            "archive.org": "Kent Beck — Test-Driven Development by Example",
            "dannorth.net": "Dan North",
            "martinfowler.com": "Martin Fowler — Are Harnesses Worth the Cost?",
        }
        label = labels.get(urlparse(url).netloc, url)
    return label, url


def inline(text: str) -> str:
    label, url = link_parts(text)
    if url:
        return f'<a href="{html.escape(url, quote=True)}" target="_blank" rel="noopener noreferrer">{html.escape(label)}</a>'
    value = html.escape(text)
    value = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", value)
    return re.sub(r"~~(.+?)~~", r"<del>\1</del>", value)


def block_html(block: str) -> str:
    lines = block.splitlines()
    main = re.match(r"^(- |\d+\. )(.*)$", lines[0])
    quote = lines[0].startswith("> ")
    if main:
        numbered = main[1] != "- "
        tag = "ol" if numbered else "ul"
        start = f' start="{main[1].split(".")[0]}"' if numbered else ""
        def children_html(children: list[str]) -> str:
            if not children:
                return ""
            level = len(children[0]) - len(children[0].lstrip())
            parts = ["<ul>"]
            i = 0
            while i < len(children):
                line = children[i]
                i += 1
                nested_lines = []
                while i < len(children) and len(children[i]) - len(children[i].lstrip()) > level:
                    nested_lines.append(children[i])
                    i += 1
                parts.append(f"<li>{inline(re.sub(r'^\s*- ', '', line))}{children_html(nested_lines)}</li>")
            return "".join(parts) + "</ul>"

        nested = children_html([line for line in lines[1:] if line.strip()])
        return f'<{tag}{start}><li>{inline(main[2])}{nested}</li></{tag}>'
    if quote:
        return f'<blockquote>{inline(block.removeprefix("> "))}</blockquote>'
    return f'<p>{inline(block).replace(chr(10), "<br>")}</p>'


# Shared coordinates: native PowerPoint shapes and offline SVG use the same graph.
NODES = {
    "1": (480, 45, "Unit 1"),
    "1.1": (100, 205, "Unit 1.1"),
    "1.2": (350, 205, "Unit 1.2"),
    "1.3": (600, 205, "Unit 1.3"),
    "1.4": (850, 205, "Unit 1.4"),
    "1.3.1": (460, 370, "Unit 1.3.1"),
    "1.3.2": (730, 370, "Unit 1.3.2"),
    "1.3.2b": (730, 520, "Unit 1.3.2b"),
}
EDGES = [("1", child) for child in ("1.1", "1.2", "1.3", "1.4")] + [("1.3", "1.3.1"), ("1.3", "1.3.2"), ("1.3.2", "1.3.2b")]


def diagram_svg() -> str:
    parts = ['<svg viewBox="0 0 1050 600" role="img" aria-label="Unit 1 decomposes into units 1.1, 1.2, 1.3 and 1.4. Unit 1.3 has two children. Unit 1.3.2 is replaced by Unit 1.3.2b.">', '<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#4aa3ff"/></marker></defs>']
    for parent, child in EDGES:
        x1, y1, _ = NODES[parent]
        x2, y2, _ = NODES[child]
        dash = ' stroke-dasharray="8 6"' if child == "1.3.2b" else ""
        parts.append(f'<line x1="{x1}" y1="{y1 + 28}" x2="{x2}" y2="{y2 - 28}" stroke="#4aa3ff" stroke-width="2" marker-end="url(#arrow)"{dash}/>')
    for key, (x, y, label) in NODES.items():
        parts.append(f'<rect x="{x-93}" y="{y-28}" width="186" height="56" rx="4" fill="#151922" stroke="#4aa3ff"/>')
        decoration = ' text-decoration="line-through"' if key == "1.3.2" else ""
        parts.append(f'<text x="{x}" y="{y+8}" text-anchor="middle" fill="#e6e6e6" font-family="Arial, sans-serif" font-size="23"{decoration}>{label}</text>')
    return "".join(parts) + "</svg>"


def textbox(slide, x, y, width, height, text, size=22, color=FG, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(text.splitlines()):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
    return box


def body_lines(item: Slide):
    for block in item.blocks:
        for j, line in enumerate(block.splitlines()):
            match = re.match(r"^(\s*)(- |\d+\. )(.*)$", line)
            raw = match[3] if match else line.removeprefix("> ")
            label, url = link_parts(clean(raw))
            indent = "  " * (len(match[1]) // 2) if match else ""
            prefix = indent + ("• " if match and match[2] == "- " else match[2] if match else "")
            yield prefix + label, url, raw, j


def body_font_size(item: Slide) -> int:
    # Fit actual wrapped text into the fixed PowerPoint content region.
    for size in range(24, 17, -1):
        try:
            font = ImageFont.truetype("arial.ttf", size * 4)
        except OSError:
            font = ImageFont.load_default(size=size * 4)
        height = 0
        for text, _, _, j in body_lines(item):
            wrapped, current = 1, ""
            for word in text.split():
                proposed = f"{current} {word}".strip()
                if current and font.getlength(proposed) > 11.54 * 288:
                    wrapped += 1
                    current = word
                else:
                    current = proposed
            height += wrapped * size * 1.25 + (12 if j == 0 else 7)
        if height <= 4.45 * 72:
            return size
    raise ValueError(f"Slide needs to be split further: {item.title}")


def export_ppt(title: str, slides: list[Slide]) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = title
    prs.core_properties.subject = "TDD presentation exported from tdd3.md"
    prs.core_properties.author = "Costin Grigore"
    for index, item in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
        textbox(slide, .7, .3, 11.9, .4, item.section, 12, MUTED)
        textbox(slide, .7, .85, 11.9, 1.05, item.title, 34, bold=True)
        if item.diagram:
            sx, sy = 11.5 / 1050, 4.65 / 600
            for parent, child in EDGES:
                x1, y1, _ = NODES[parent]
                x2, y2, _ = NODES[child]
                edge = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(.8 + x1*sx), Inches(2 + (y1+28)*sy), Inches(.8 + x2*sx), Inches(2 + (y2-28)*sy))
                edge.line.color.rgb = RGBColor.from_string(BLUE)
                edge.line.width = Pt(1.5)
                arrow = OxmlElement("a:tailEnd")
                arrow.set("type", "triangle")
                edge.line._get_or_add_ln().append(arrow)
                if child == "1.3.2b":
                    edge.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            for key, (x, y, label) in NODES.items():
                shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.8+(x-93)*sx), Inches(2+(y-28)*sy), Inches(186*sx), Inches(56*sy))
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(PANEL)
                shape.line.color.rgb = RGBColor.from_string(BLUE)
                shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = shape.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                p.text = label
                p.font.name = "Arial"
                p.font.size = Pt(17)
                p.font.color.rgb = RGBColor.from_string(FG)
                if key == "1.3.2":
                    for run in p.runs:
                        run._r.get_or_add_rPr().set("strike", "sngStrike")
        else:
            box = textbox(slide, .85, 2.05, 11.6, 4.55, "", 23)
            tf = box.text_frame
            size = body_font_size(item)
            first = True
            for text, url, raw, j in body_lines(item):
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.text = text
                p.font.name = "Arial"
                p.font.size = Pt(size)
                p.font.color.rgb = RGBColor.from_string(BLUE if url else FG)
                p.line_spacing = 1.2
                p.space_before = Pt(9 if j == 0 else 4)
                p.space_after = Pt(3)
                if "~~" in raw:
                    for run in p.runs:
                        run._r.get_or_add_rPr().set("strike", "sngStrike")
                if url:
                    p.runs[0].hyperlink.address = url
        textbox(slide, .7, 7.0, 10.5, .3, "TDD · Where did it all go wrong?", 10, MUTED)
        textbox(slide, 11.6, 7.0, 1.0, .3, f"{index+1} / {len(slides)}", 10, MUTED)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.45), int(prs.slide_width*(index+1)/len(slides)), Inches(.05))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string(BLUE)
        bar.line.fill.background()
        slide.notes_slide.notes_text_frame.text = f"Source: tdd3.md\n\n{item.source}"
    prs.save(ROOT / "tdd-presentation.pptx")


def export_html(title: str, slides: list[Slide]) -> None:
    # Retain the reference deck's layout, colors, typography and interaction model.
    reference = (ROOT / "2026-07-01--tdd" / "tdd-present1.html").read_text(encoding="utf-8")
    css = re.search(r"<style>(.*?)</style>", reference, re.S)[1]
    extra_css = """
  * { box-sizing: border-box; }
  #sidebar { flex: 0 0 260px; }
  #main { min-width: 0; min-height: 0; }
  #content { min-height: 0; }
  #nav { display: grid; }
  button.slide-link { text-align: left; border: 0; border-bottom: 1px solid #222733; background: transparent; }
  button.slide-link:hover { background: #1f2633; }
  button.active-link { background: #27324a; }
  .slide { max-width: 1100px; width: 100%; margin-inline: auto; font-size: clamp(20px, 1.9vw, 26px); line-height: 1.35; overflow-wrap: anywhere; }
  .slide h2 { font-size: clamp(30px, 3.2vw, 44px); line-height: 1.15; margin: 12px 0 22px; }
  .slide ul, .slide ol { padding-left: 1.4em; margin-block: .3em; }
  .slide li > ul { color: #c4cad5; }
  .slide > ul > li > ul, .slide > ol > li > ul { font-size: .86em; }
  .slide p { margin-block: .6em; }
  .slide blockquote { border-left: 3px solid #4aa3ff; margin: 1em 0; padding-left: 1em; }
  .slide svg { display: block; max-height: calc(100dvh - 330px); width: 100%; min-height: 180px; }
  a { color: #4aa3ff; }
  .tag { font-size: 12px; }
  #controls { align-items: center; gap: 12px; }
  #counter { font-size: 13px; color: #aaa; }
  button:disabled { opacity: .4; cursor: default; }
  button:focus-visible, a:focus-visible { outline: 2px solid #4aa3ff; outline-offset: 3px; }
  @media (max-width: 760px) {
    body { height: 100dvh; }
    #sidebar { flex-basis: 150px; width: 150px; }
    button.slide-link { font-size: 11px; padding: 9px 8px; }
    #content { padding: 20px 16px; }
    #header h1 { font-size: 14px; }
    #controls { gap: 6px; }
    #controls button { font-size: 12px; padding: 6px 7px; }
    #counter { font-size: 11px; }
    .slide { font-size: 18px; }
    .slide h2 { font-size: 28px; }
  }
  @media (prefers-reduced-motion: reduce) { .slide { animation: none; } }
  @media print {
    body { display: block; height: auto; overflow: visible; background: white; color: black; }
    #sidebar, #header, #controls { display: none; }
    #main, #content { display: block; overflow: visible; padding: 0; }
    .slide { display: block !important; break-after: page; padding: 20px; font-size: 20px; animation: none; }
    .slide:last-child { break-after: auto; }
    .tag, a, .slide li > ul { color: inherit; }
    .slide svg { max-height: none; }
    @page { size: landscape; margin: 12mm; }
  }
"""
    sections = []
    for i, item in enumerate(slides):
        body = diagram_svg() if item.diagram else "".join(block_html(block) for block in item.blocks)
        sections.append(f'<section class="slide" id="slide-{i+1}" aria-labelledby="title-{i+1}" hidden><div class="tag">{html.escape(item.section)}</div><h2 id="title-{i+1}">{html.escape(item.title)}</h2>{body}</section>')
    titles = json.dumps([item.title for item in slides], ensure_ascii=False).replace("<", "\\u003c")
    page = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{html.escape(title)}</title>
<style>{css}{extra_css}</style>
</head>
<body>
<aside id="sidebar" aria-label="Slides"><h2>Slides</h2><nav id="nav" aria-label="Slide navigation"></nav></aside>
<main id="main">
  <header id="header"><h1>{html.escape(title)}</h1><small>Presenter: Costin Grigore · 2026-10-07</small><div id="progressBar" role="progressbar" aria-label="Presentation progress" aria-valuemin="1" aria-valuemax="{len(slides)}"><div id="progress"></div></div></header>
  <div id="content">{''.join(sections)}</div>
  <footer id="controls"><button id="prev" type="button">⬅ Prev</button><span id="counter" aria-live="polite"></span><button id="next" type="button">Next ➡</button></footer>
</main>
<script>
const titles = {titles};
const slides = [...document.querySelectorAll('.slide')];
const nav = document.getElementById('nav');
let index = 0;
titles.forEach((title, i) => {{
  const button = document.createElement('button');
  button.type = 'button'; button.className = 'slide-link';
  button.textContent = `${{i + 1}}. ${{title}}`;
  button.addEventListener('click', () => go(i));
  nav.appendChild(button);
}});
const links = [...nav.children];
function render() {{
  slides.forEach((slide, i) => {{ slide.hidden = i !== index; slide.classList.toggle('active', i === index); }});
  links.forEach((link, i) => {{
    link.classList.toggle('active-link', i === index);
    if (i === index) link.setAttribute('aria-current', 'step'); else link.removeAttribute('aria-current');
  }});
  document.getElementById('counter').textContent = `${{index + 1}} / ${{slides.length}}`;
  document.getElementById('progress').style.width = `${{(index + 1) / slides.length * 100}}%`;
  document.getElementById('progressBar').setAttribute('aria-valuenow', index + 1);
  document.getElementById('prev').disabled = index === 0;
  document.getElementById('next').disabled = index === slides.length - 1;
  document.getElementById('content').scrollTop = 0;
  links[index].scrollIntoView({{ block: 'nearest' }});
}}
function go(i) {{
  index = Math.max(0, Math.min(slides.length - 1, i));
  history.replaceState(null, '', `#slide-${{index + 1}}`);
  render();
}}
function readHash() {{
  const match = /^#slide-(\\d+)$/.exec(location.hash);
  index = match ? Math.max(0, Math.min(slides.length - 1, Number(match[1]) - 1)) : 0;
  render();
}}
document.getElementById('prev').addEventListener('click', () => go(index - 1));
document.getElementById('next').addEventListener('click', () => go(index + 1));
document.addEventListener('keydown', event => {{
  if (event.altKey || event.ctrlKey || event.metaKey || /INPUT|TEXTAREA|SELECT/.test(event.target.tagName)) return;
  if (['ArrowRight', 'PageDown'].includes(event.key)) {{ event.preventDefault(); go(index + 1); }}
  if (['ArrowLeft', 'PageUp'].includes(event.key)) {{ event.preventDefault(); go(index - 1); }}
  if (event.key === 'Home') {{ event.preventDefault(); go(0); }}
  if (event.key === 'End') {{ event.preventDefault(); go(slides.length - 1); }}
}});
window.addEventListener('hashchange', readHash);
readHash();
</script>
</body>
</html>
"""
    (ROOT / "tdd-presentation.html").write_text(page, encoding="utf-8")


if __name__ == "__main__":
    deck_title, deck_slides = read_slides()
    export_ppt(deck_title, deck_slides)
    export_html(deck_title, deck_slides)
    print(f"Exported {len(deck_slides)} slides to tdd-presentation.pptx and tdd-presentation.html")
